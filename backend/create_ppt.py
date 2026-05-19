from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    
    # Set slide dimensions to 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Helper function to add a title slide
    def add_title_slide(title_text, subtitle_text):
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        # Change background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(8, 12, 10)  # Dark background
        
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = title_text
        subtitle.text = subtitle_text
        
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.name = 'Arial'
        
        # Style specifically the title
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(64)
        title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(61, 220, 151) # Green accent
        return slide

    # Helper function to add a content slide
    def add_content_slide(title_text, content_points):
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(17, 23, 20)
        
        title = slide.shapes.title
        body = slide.placeholders[1]
        
        title.text = title_text
        title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(61, 220, 151)
        
        tf = body.text_frame
        tf.clear() # clear default paragraphs
        
        for point in content_points:
            p = tf.add_paragraph()
            p.text = point
            p.font.size = Pt(28)
            p.font.color.rgb = RGBColor(232, 236, 234)
            p.level = 0
            # Space after paragraph
            p.space_after = Pt(14)
            
        return slide

    # Slide 1: Title
    add_title_slide(
        "Citizen Health AI", 
        "AI-Powered Multilingual Voice Assistant for Remote Healthcare\nBuilt with Sarvam AI, RAG, and MCP"
    )

    # Slide 2: Business Problem
    add_content_slide(
        "The Business Problem",
        [
            "7-8 crore citizens in Tamil Nadu and Karnataka face fragmented healthcare access.",
            "Significant gaps exist in tribal and remote PHCs.",
            "Staff spend significant time answering repetitive queries across hospitals and schemes.",
            "Language barriers prevent citizens from navigating the health system effectively.",
            "Currently, there is zero AI-assisted proactive outreach."
        ]
    )

    # Slide 3: The Solution
    add_content_slide(
        "Our Solution: Citizen Health AI",
        [
            "A fully autonomous, pure voice-only AI agent.",
            "Speaks fluent Tamil, Kannada, Hindi, and English natively via Sarvam AI.",
            "Solves 5 core PHC interactions instantly and autonomously.",
            "Reduces staff workload by answering queries and booking appointments.",
            "Operates via a beautiful, accessible 'Orb' web interface."
        ]
    )

    # Slide 4: Core Features
    add_content_slide(
        "The 5 Core Modules (Implemented)",
        [
            "1. Hospital Navigation: Guides patients to the right floor/room.",
            "2. Eligibility Checking: Instantly checks Ayushman Bharat & CMCHIS eligibility.",
            "3. Appointment Booking: Books follow-ups autonomously with calendar slot-filling.",
            "4. Maternal Health Reminders: Schedules ANC checkup reminders.",
            "5. NHM Programme Queries: Answers JSY/RKSK queries accurately without hallucinations."
        ]
    )

    # Slide 5: System Architecture
    add_content_slide(
        "High-Level Architecture",
        [
            "Frontend: Next.js 14, React, Web Audio API (Voice Activity Detection).",
            "Backend: FastAPI serverless deployment on Vercel.",
            "Speech Engine: Sarvam AI (STT & TTS) for low-latency Indic languages.",
            "LLM Engine: Sarvam-30b / Sarvam-m for native language reasoning.",
            "Knowledge Base: Supabase PostgreSQL with pgvector for RAG.",
            "Orchestration: Advanced State Machine (Verification -> Intent -> Action)."
        ]
    )

    # Slide 6: Identity Verification Flow
    add_content_slide(
        "Seamless Identity Verification",
        [
            "1. Wake Word Detection: Browser detects 'Namaste', 'Vanakkam', 'Hello'.",
            "2. Agent Introduction: Greets and asks for a 10-digit mobile or ABHA ID.",
            "3. LLM Extraction: Parses natural speech to extract digits.",
            "4. DB Verification: Checks Supabase 'citizens' table.",
            "5. State Update: Upgrades session to 'Verified' and personalizes greeting.",
            "6. Fallback: Guest mode for unregistered citizens."
        ]
    )

    # Slide 7: RAG & MCP Integration
    add_content_slide(
        "RAG & MCP (Model Context Protocol)",
        [
            "RAG (Retrieval-Augmented Generation):",
            "  - Prevents hallucinations on healthcare schemes.",
            "  - Injects official NHM documents from pgvector into the prompt context.",
            "MCP (Model Context Protocol):",
            "  - LLM outputs strict JSON indicating tool calls (e.g., book_appointment).",
            "  - Backend executes DB queries and feeds results back to LLM.",
            "  - Ensures reliable, deterministic database mutations."
        ]
    )

    # Slide 8: Safety & Triage
    add_content_slide(
        "Safety-First: Symptom Triage Pipeline",
        [
            "Keyword interceptor runs BEFORE the LLM prompt.",
            "Emergency Symptoms (e.g., chest pain, heavy bleeding):",
            "  - Immediately overrides standard response.",
            "  - Instructs citizen to 'Call 108 immediately'.",
            "  - Flags interaction with 'needs_ambulance = true'.",
            "Non-Emergency: Classifies as High, Medium, or Low severity.",
            "Guides patients safely to the right level of care."
        ]
    )

    # Slide 9: Production Readiness
    add_content_slide(
        "Production-Grade & Tested",
        [
            "Fully automated test suite: 37/37 passing tests.",
            "LLM Intelligence Probe: Verifies JSON compliance & slot-filling strictly.",
            "Graceful Degradation: Handles silent audio & API limits smoothly.",
            "Analytics Dashboard: Tracks intents, languages, and emergency counts.",
            "Deployment: Ready on Vercel (Backend + Frontend separated)."
        ]
    )

    prs.save('Citizen_Health_AI_Presentation.pptx')
    print("Presentation created successfully as Citizen_Health_AI_Presentation.pptx")

if __name__ == '__main__':
    create_presentation()
